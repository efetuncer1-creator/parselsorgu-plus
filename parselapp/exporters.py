"""KML ve Excel çıktıları.

KML: Google Earth'te açıldığında Alan Yöneticisi'ndekine benzer bilgi tablosu gösterir;
ayrıca ExtendedData/Data alanları taşıdığı için bu uygulamaya geri aktarılabilir.
Excel: "Parseller" (TKGM bilgileri) ve "Koordinatlar" (ITRF96 / ED50 3° köşe noktaları).
"""

import io
import re
import zipfile
from collections import OrderedDict
from datetime import datetime
from html import escape
from xml.sax.saxutils import escape as xml_escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from shapely.geometry import Polygon

from . import coords

# ---- ortak ---------------------------------------------------------------------


def ad_parsel(p):
    return f"{p['ada']}/{p['parsel']}"


def merkez(p):
    """Etiket/navigasyon noktası: en büyük poligonun içinde kalan bir nokta (boylam, enlem)."""
    pol = max((Polygon(pg[0], pg[1:]) for pg in p["geometri"]), key=lambda g: g.area)
    nokta = pol.representative_point()
    return nokta.x, nokta.y


def harita_linki(p):
    lon, lat = merkez(p)
    return f"https://www.google.com/maps/search/?api=1&query={lat:.7f},{lon:.7f}"


def _guvenli(s):
    return re.sub(r'[\\/:*?"<>|]+', "-", s).strip(" .-") or "parsel"


def dosya_adi(parseller, uzanti):
    if len(parseller) == 1:
        p = parseller[0]
        ad = f"{p['il']}-{p['ilce']}-{p['mahalle']}-{p['ada']}-{p['parsel']}"
    else:
        yerler = {(p["il"], p["ilce"], p["mahalle"]) for p in parseller}
        if len(yerler) == 1:
            il, ilce, mah = yerler.pop()
            ad = f"{il}-{ilce}-{mah}-{len(parseller)} parsel"
        else:
            ad = f"Parseller-{len(parseller)} parsel-{datetime.now():%Y-%m-%d}"
    return _guvenli(ad) + uzanti


# ---- KML ---------------------------------------------------------------------------

_KML_BAS = """<?xml version="1.0" encoding="UTF-8"?>
<kml xmlns="http://www.opengis.net/kml/2.2">
<Document>
  <name>{ad}</name>
  <Style id="parsel">
    <LineStyle><color>ff0000ff</color><width>3</width></LineStyle>
    <PolyStyle><color>33ffffff</color></PolyStyle>
    <BalloonStyle><text>$[description]</text></BalloonStyle>
  </Style>
  <Style id="sinir">
    <LineStyle><color>ffff7b2b</color><width>3</width></LineStyle>
    <PolyStyle><fill>0</fill></PolyStyle>
  </Style>
  <Style id="etiket">
    <IconStyle><scale>0</scale></IconStyle>
    <LabelStyle><color>ff00aaff</color><scale>0.9</scale></LabelStyle>
  </Style>
"""

_BILGI_SATIRLARI = (
    ("İl", "il"), ("İlçe", "ilce"), ("Mahalle", "mahalle"), ("Ada", "ada"), ("Parsel", "parsel"),
    ("Tapu Alanı", "alan"), ("Nitelik", "nitelik"), ("Mevkii", "mevkii"), ("Pafta", "pafta"),
    ("Zemin", "zemin"),
)

# Geri aktarımda importers._ETIKETLER ile eşleşen adlar
_DATA_ADLARI = (
    ("İl", "il"), ("İlçe", "ilce"), ("Mahalle", "mahalle"), ("Ada", "ada"), ("ParselNo", "parsel"),
    ("Alan", "alan"), ("Nitelik", "nitelik"), ("Mevkii", "mevkii"), ("Pafta", "pafta"),
    ("Zemin", "zemin"), ("mahalle_id", "mahalle_id"),
)


def _aciklama(p):
    satirlar = "".join(
        f"<tr><td><b>{etiket}</b></td><td>{escape(str(p.get(alan) or ''))}</td></tr>"
        for etiket, alan in _BILGI_SATIRLARI)
    satirlar += f'<tr><td><b>Navigasyon</b></td><td><a href="{escape(harita_linki(p))}">Google Haritalar</a></td></tr>'
    return ('<table style="font-family:Arial,sans-serif;border-collapse:collapse" border="1" cellpadding="5">'
            f"{satirlar}</table>")


def _halka(ring):
    return " ".join(f"{lon:.7f},{lat:.7f},0" for lon, lat in ring)


def _placemark(p):
    data = "".join(
        f'<Data name="{ad}"><value>{xml_escape(str(p.get(alan)))}</value></Data>'
        for ad, alan in _DATA_ADLARI if p.get(alan) not in (None, ""))
    poligonlar = []
    for pol in p["geometri"]:
        ic = "".join(f"<innerBoundaryIs><LinearRing><coordinates>{_halka(r)}</coordinates></LinearRing></innerBoundaryIs>"
                     for r in pol[1:])
        poligonlar.append(f"<Polygon><outerBoundaryIs><LinearRing><coordinates>{_halka(pol[0])}"
                          f"</coordinates></LinearRing></outerBoundaryIs>{ic}</Polygon>")
    geom = poligonlar[0] if len(poligonlar) == 1 else f"<MultiGeometry>{''.join(poligonlar)}</MultiGeometry>"
    lon, lat = merkez(p)
    ad = xml_escape(ad_parsel(p))
    return (f"    <Placemark><name>{ad}</name><styleUrl>#parsel</styleUrl>"
            f"<description><![CDATA[{_aciklama(p)}]]></description>"
            f"<ExtendedData>{data}</ExtendedData>{geom}</Placemark>\n"
            f"    <Placemark><name>{ad}</name><styleUrl>#etiket</styleUrl>"
            f"<Point><coordinates>{lon:.7f},{lat:.7f},0</coordinates></Point></Placemark>\n")


def _sinir_klasoru(ref):
    """Referans alanının (açık katmanları) KML klasörü."""
    gizli = set(ref.get("gizli") or [])
    hesap = ref.get("hesap_katmani")
    yer = []
    for f in ref["geojson"]["features"]:
        props = f.get("properties") or {}
        katman = props.get("katman", "")
        if (katman != hesap) if hesap is not None else (katman in gizli):
            continue
        g = f["geometry"]
        if g["type"] == "Polygon":
            ic = "".join(f"<innerBoundaryIs><LinearRing><coordinates>{_halka(r)}</coordinates></LinearRing></innerBoundaryIs>"
                         for r in g["coordinates"][1:])
            geom = (f"<Polygon><outerBoundaryIs><LinearRing><coordinates>{_halka(g['coordinates'][0])}"
                    f"</coordinates></LinearRing></outerBoundaryIs>{ic}</Polygon>")
        elif g["type"] == "LineString":
            geom = f"<LineString><coordinates>{_halka(g['coordinates'])}</coordinates></LineString>"
        elif g["type"] == "Point":
            lon, lat = g["coordinates"][:2]
            geom = f"<Point><coordinates>{lon:.7f},{lat:.7f},0</coordinates></Point>"
        else:
            continue
        yer.append(f"    <Placemark><name>{xml_escape(props.get('ad') or '')}</name>"
                   f"<styleUrl>#sinir</styleUrl>{geom}</Placemark>\n")
    return f"  <Folder><name>{xml_escape(ref['ad'])} (alan sınırı)</name>\n{''.join(yer)}  </Folder>\n"


def kml(parseller, belge_adi=None, referans=None):
    """Tek KML; parseller mahalle klasörlerine ayrılır. referans verilirse sınırı da eklenir."""
    gruplar = OrderedDict()
    for p in parseller:
        gruplar.setdefault(f"{p['il']} / {p['ilce']} / {p['mahalle']}", []).append(p)
    parcalar = [_KML_BAS.format(ad=xml_escape(belge_adi or dosya_adi(parseller, "")))]
    if referans:
        parcalar.append(_sinir_klasoru(referans))
    for grup, liste in gruplar.items():
        parcalar.append(f"  <Folder><name>{xml_escape(grup)}</name>\n")
        parcalar.extend(_placemark(p) for p in liste)
        parcalar.append("  </Folder>\n")
    parcalar.append("</Document>\n</kml>\n")
    return "".join(parcalar).encode("utf-8")


def kml_ayri_zip(parseller):
    """Her parsel ayrı KML, hepsi bir zip içinde."""
    tampon = io.BytesIO()
    with zipfile.ZipFile(tampon, "w", zipfile.ZIP_DEFLATED) as z:
        kullanilan = set()
        for p in parseller:
            ad = dosya_adi([p], ".kml")
            if ad in kullanilan:
                ad = ad[:-4] + f"-{len(kullanilan)}.kml"
            kullanilan.add(ad)
            z.writestr(ad, kml([p]))
    return tampon.getvalue()


# ---- Excel ---------------------------------------------------------------------------

_BASLIK_STIL = dict(font=Font(bold=True, color="FFFFFF"), fill=PatternFill("solid", fgColor="2F5D50"),
                    alignment=Alignment(vertical="center", wrap_text=True))


def _baslik_yaz(ws, satir, basliklar):
    for i, b in enumerate(basliklar, 1):
        c = ws.cell(row=satir, column=i, value=b)
        for k, v in _BASLIK_STIL.items():
            setattr(c, k, v)


def _genislik(ws, genislikler):
    for i, g in enumerate(genislikler, 1):
        ws.column_dimensions[get_column_letter(i)].width = g


def dilim_belirle(parseller, dilim="otomatik"):
    """Tüm çıktı tek dilimde olmalı; otomatikte parsellerin ortalama boylamı kullanılır."""
    if dilim not in (None, "", "otomatik"):
        dom = int(dilim)
        if dom not in coords.DILIMLER:
            raise ValueError(f"Geçersiz dilim: {dilim}")
        return dom
    boylamlar = [merkez(p)[0] for p in parseller]
    return coords.dilim_sec(sum(boylamlar) / len(boylamlar))


def excel(parseller, sistem="itrf96", dilim="otomatik", alan_ici=None, alan_adi=None):
    """alan_ici: {key: (m², oran)} verilirse 'Alan İçi' sütunları eklenir."""
    if sistem not in coords.SISTEMLER:
        raise ValueError(f"Geçersiz koordinat sistemi: {sistem}")
    dom = dilim_belirle(parseller, dilim)
    wb = Workbook()

    # Parseller
    ws = wb.active
    ws.title = "Parseller"
    basliklar = ["Sıra", "İl", "İlçe", "Mahalle", "Ada", "Parsel", "Tapu Alanı (m²)", "Nitelik",
                 "Mevkii", "Pafta", "Zemin", "Navigasyon", "Kaynak", "Veri Tarihi"]
    if alan_ici is not None:
        basliklar += [f"{alan_adi or 'Alan'} içinde (m², yaklaşık)", "Alan içi oran"]
    _baslik_yaz(ws, 1, basliklar)
    for i, p in enumerate(parseller, 1):
        link = harita_linki(p)
        satir = [i, p["il"], p["ilce"], p["mahalle"], _sayi(p["ada"]), _sayi(p["parsel"]), p["alan_m2"],
                 p["nitelik"], p["mevkii"], p["pafta"], p["zemin"], "Google Haritalar",
                 "TKGM" if p["kaynak"] == "tkgm" else p["kaynak"].removeprefix("dosya:"),
                 p["kaynak_tarih"][:10]]
        if alan_ici is not None:
            ic, oran = alan_ici.get(p["key"], (0.0, 0.0))
            satir += [round(ic, 2), round(oran, 4)]
        ws.append(satir)
        ws.cell(row=i + 1, column=7).number_format = "#,##0.00"
        if alan_ici is not None:
            ws.cell(row=i + 1, column=15).number_format = "#,##0.00"
            ws.cell(row=i + 1, column=16).number_format = "0.0%"
        c = ws.cell(row=i + 1, column=12)
        c.hyperlink = link
        c.font = Font(color="0563C1", underline="single")
    son = len(parseller) + 1
    ws.cell(row=son + 1, column=6, value="Toplam").font = Font(bold=True)
    t = ws.cell(row=son + 1, column=7, value=f"=SUM(G2:G{son})")
    t.number_format = "#,##0.00"
    t.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(basliklar))}{son}"
    _genislik(ws, [6, 12, 14, 18, 8, 8, 16, 18, 18, 12, 14, 18, 30, 12, 22, 12])
    if alan_ici is not None:
        t2 = ws.cell(row=son + 1, column=15, value=f"=SUM(O2:O{son})")
        t2.number_format = "#,##0.00"
        t2.font = Font(bold=True)

    # Koordinatlar
    wk = wb.create_sheet("Koordinatlar")
    sistem_adi = coords.SISTEMLER[sistem]["ad"]
    wk["A1"] = f"Koordinat sistemi: {sistem_adi}, DOM {dom}° (EPSG:{coords.epsg_kodu(sistem, dom)})"
    wk["A1"].font = Font(bold=True)
    wk["A2"] = "Y = sağa (doğu), X = yukarı (kuzey). Kaynak: TKGM, WGS84 enlem/boylam."
    _baslik_yaz(wk, 4, ["Mahalle", "Ada/Parsel", "Halka", "Nokta No", "Y (Sağa)", "X (Yukarı)", "Enlem", "Boylam"])
    for p in parseller:
        for pi, pol in enumerate(p["geometri"], 1):
            for hi, ring in enumerate(pol):
                noktalar = ring[:-1] if len(ring) > 1 and ring[0] == ring[-1] else ring
                yx = coords.donustur_halka(noktalar, sistem, dom)
                halka = ("Dış" if hi == 0 else f"İç {hi}") + (f" (Poligon {pi})" if len(p["geometri"]) > 1 else "")
                for ni, ((lon, lat), (y, x)) in enumerate(zip(noktalar, yx), 1):
                    wk.append([p["mahalle"], ad_parsel(p), halka, ni, round(y, 2), round(x, 2), lat, lon])
                    r = wk.max_row
                    wk.cell(row=r, column=5).number_format = "0.00"
                    wk.cell(row=r, column=6).number_format = "0.00"
    wk.freeze_panes = "A5"
    _genislik(wk, [18, 14, 12, 10, 14, 14, 13, 13])

    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue(), dom


def _sayi(s):
    try:
        return int(s)
    except (TypeError, ValueError):
        return s
