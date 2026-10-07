"""Alandan parsel tespiti testleri; TKGM'ye bağlanmaz."""

import io
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from shapely.geometry import box

from parselapp import analiz, coords, exporters, referans
from parselapp.importers import kml_oku
from parselapp.store import Store
from parselapp.tkgm import LimitAsildi, Tkgm

from .test_parselapp import _Cevap, _SahteOturum

DOM = 30
X0, Y0, KENAR = 500000.0, 4560000.0, 60.0


def hucre(i, j):
    """Izgaranın (i, j) hücresi, coğrafi koordinatlarla dış halka."""
    g = coords.cografi(box(X0 + i * KENAR, Y0 + j * KENAR, X0 + (i + 1) * KENAR, Y0 + (j + 1) * KENAR), DOM)
    return [[round(x, 7), round(y, 7)] for x, y in g.exterior.coords]


def kayit(i, j, alan="3.600,00", nitelik="Tarla"):
    return {"il": "Örnekil", "ilce": "Örnekilçe", "mahalle": "Izgara", "ada": "1", "parsel": str(i * 3 + j + 1),
            "mahalle_id": 500, "alan": alan, "nitelik": nitelik, "geometri": [[hucre(i, j)]]}


def tkgm_cevabi(i, j):
    return {"type": "Feature", "geometry": {"type": "Polygon", "coordinates": [hucre(i, j)]},
            "properties": {"ilAd": "Örnekil", "ilceAd": "Örnekilçe", "mahalleAd": "Izgara", "adaNo": "1",
                           "parselNo": str(i * 3 + j + 1), "mahalleId": 500, "alan": "3.600,00", "nitelik": "Bağ"}}


class AlanTest(unittest.TestCase):
    def setUp(self):
        import parselapp.tkgm as t
        self._bekleme, t.BEKLEME_SN = t.BEKLEME_SN, 0
        self.klasor = tempfile.TemporaryDirectory()
        self.store = Store(Path(self.klasor.name) / "t.db")
        # Ortadaki (1, 1) hücre hariç 8 parsel kayıtlı
        for i in range(3):
            for j in range(3):
                if (i, j) != (1, 1):
                    self.store.kaydet(kayit(i, j, nitelik="Orman" if i == 0 else "Tarla"), "dosya:izgara.kml")
        dis = coords.cografi(box(X0, Y0, X0 + 3 * KENAR, Y0 + 3 * KENAR), DOM)
        fc = {"type": "FeatureCollection", "features": [
            {"type": "Feature", "geometry": json.loads(json.dumps(dis.__geo_interface__)),
             "properties": {"ad": "Izgara", "katman": ""}}]}
        self.ref_id = self.store.referans_ekle("Izgara OSB", "#2b7bff", fc)

    def tearDown(self):
        import parselapp.tkgm as t
        t.BEKLEME_SN = self._bekleme
        self.store._db.close()
        self.klasor.cleanup()

    def ref(self):
        return self.store.referans(self.ref_id)

    def oturum(self):
        return [p["key"] for p in self.store.listele()]

    def test_ozet(self):
        o = analiz.ozet(self.ref(), self.store.listele())
        self.assertAlmostEqual(o["alan_m2"], 9 * 3600, delta=5)
        self.assertEqual(o["icinde"]["adet"], 8)
        self.assertAlmostEqual(o["icinde"]["tapu_m2"], 8 * 3600)
        self.assertAlmostEqual(o["kaplanan_oran"], 8 / 9, places=3)
        self.assertEqual(o["bosluk"]["adet"], 1)
        self.assertAlmostEqual(o["bosluk"]["m2"], 3600, delta=5)
        self.assertEqual({n["nitelik"]: n["adet"] for n in o["nitelikler"]}, {"Tarla": 5, "Orman": 3})

    def test_tarama_boslugu_doldurur(self):
        oturum = _SahteOturum([_Cevap(200, tkgm_cevabi(1, 1))])
        t = Tkgm(self.store, oturum)
        s1 = analiz.tara_adim(self.store, t, self.ref(), self.oturum())
        self.assertEqual(s1["durum"], "bulundu")
        self.assertTrue(s1["sorgu"])
        self.assertEqual(s1["parsel"]["parsel"], "5")
        lon, lat = s1["nokta"]
        self.assertIn(f"/parsel/{lat:.6f}/{lon:.6f}/", oturum.istekler[0])
        s2 = analiz.tara_adim(self.store, t, self.ref(), self.oturum())
        self.assertEqual(s2["durum"], "bitti")
        self.assertEqual(len(oturum.istekler), 1)  # ikinci adımda sorgu atılmadı

    def test_parsel_cikmayan_nokta_bir_daha_sorulmaz(self):
        oturum = _SahteOturum([_Cevap(200, "null")])
        t = Tkgm(self.store, oturum)
        self.assertEqual(analiz.tara_adim(self.store, t, self.ref(), self.oturum())["durum"], "bos")
        self.assertEqual(len(self.ref()["taranan"]), 1)
        self.assertEqual(analiz.tara_adim(self.store, t, self.ref(), self.oturum())["durum"], "bitti")
        self.assertEqual(len(oturum.istekler), 1)
        self.store.referans_tarama_sifirla(self.ref_id)
        self.assertEqual(self.ref()["taranan"], [])

    def test_limit_taramayi_durdurur(self):
        t = Tkgm(self.store, _SahteOturum([_Cevap(403, '"Günlük sorgu limitini aştınız."')]))
        with self.assertRaises(LimitAsildi):
            analiz.tara_adim(self.store, t, self.ref(), self.oturum())
        self.assertEqual(self.ref()["taranan"], [])  # limit "boş nokta" sayılmaz

    def test_onbellekteki_parsel_sorgu_harcamadan_oturuma_gelir(self):
        oturum = self.oturum()                            # ekranda 8 parsel
        self.store.kaydet(kayit(1, 1), "tkgm")             # ortadaki önceki bir oturumdan önbellekte
        t = Tkgm(self.store, _SahteOturum([]))             # TKGM'ye istek atılırsa hata verir
        s = analiz.tara_adim(self.store, t, self.ref(), oturum)
        self.assertEqual((s["durum"], s["sorgu"], s["parsel"]["parsel"]), ("bulundu", False, "5"))
        # Oturumdaki parseller boşken özet de boş: önbellek ekrana karışmaz
        self.assertEqual(analiz.ozet(self.ref(), self.store.listele([]))["icinde"]["adet"], 0)

    def test_ince_seritler_bosluk_sayilmaz(self):
        # Ortadaki parseli 1 m içeriden kaydet: kenarlarda 1 m'lik şeritler kalır
        g = coords.cografi(box(X0 + KENAR + 1, Y0 + KENAR + 1, X0 + 2 * KENAR - 1, Y0 + 2 * KENAR - 1), DOM)
        k = kayit(1, 1)
        k["geometri"] = [[[list(c) for c in g.exterior.coords]]]
        self.store.kaydet(k, "dosya:x.kml")
        self.assertEqual(analiz.ozet(self.ref(), self.store.listele())["bosluk"]["adet"], 0)

    def test_alan_ici_excel_ve_kml_siniri(self):
        parseller = self.store.listele()
        alan_ici = analiz.alan_ici_m2(self.ref(), parseller)
        veri, _ = exporters.excel(parseller, alan_ici=alan_ici, alan_adi="Izgara OSB")
        ws = load_workbook(io.BytesIO(veri))["Parseller"]
        self.assertEqual(ws["O1"].value, "Izgara OSB içinde (m², yaklaşık)")
        self.assertAlmostEqual(ws["O2"].value, 3600, delta=5)
        self.assertAlmostEqual(ws["P2"].value, 1.0, places=3)
        kml = exporters.kml(parseller, referans=self.ref()).decode("utf-8")
        self.assertIn("Izgara OSB (alan sınırı)", kml)
        self.assertEqual(len(kml_oku(kml.encode("utf-8"))), 8)  # sınır parsel olarak geri okunmaz


def _ozellik(katman, geom):
    return {"type": "Feature", "geometry": geom, "properties": {"ad": "", "katman": katman}}


def _kare(x, y, d):
    return {"type": "Polygon", "coordinates": [[[x, y], [x + d, y], [x + d, y + d], [x, y + d], [x, y]]]}


class SinirKatmaniTest(unittest.TestCase):
    def setUp(self):
        # Netcad OSB projesine benzer: pafta indeksi, plan adaları, sınır
        self.fc = {"type": "FeatureCollection", "features": [
            _ozellik("PINDEX_100000", _kare(30.0, 40.0, 1.0)),
            *[_ozellik("PL_A", _kare(30.40 + i * 0.001, 40.40, 0.0008)) for i in range(5)],
            _ozellik("SNR_HESAP_OSB", _kare(30.39, 40.39, 0.02)),
            _ozellik("SNR_OSB", _kare(30.395, 40.395, 0.012)),
            _ozellik("SM_FONKSIYON", {"type": "Point", "coordinates": [30.401, 40.401]}),
        ]}

    def test_oneri(self):
        self.assertEqual(referans.sinir_katmani_oner(self.fc), "SNR_OSB")

    def test_tek_katmanli_dosyada_oneri_yok(self):
        fc = {"type": "FeatureCollection", "features": [_ozellik("", _kare(30, 40, 0.01))]}
        self.assertIsNone(referans.sinir_katmani_oner(fc))

    def test_ipucu_yoksa_oneri_yok(self):
        fc = {"type": "FeatureCollection", "features": [_ozellik("A", _kare(30, 40, 0.01)), _ozellik("B", _kare(31, 40, 0.01))]}
        self.assertIsNone(referans.sinir_katmani_oner(fc))

    def test_hesap_katmani_alani_belirler(self):
        tum = referans.alan(self.fc)
        sinir = referans.alan(self.fc, hesap_katmani="SNR_OSB")
        self.assertAlmostEqual(tum.area, 1.0, places=3)            # pafta indeksi her şeyi yutar
        self.assertAlmostEqual(sinir.area, 0.012 ** 2, places=8)
        # Hesap katmanı gizli katmanlar listesinde olsa da kullanılır
        self.assertAlmostEqual(referans.alan(self.fc, ["SNR_OSB"], "SNR_OSB").area, 0.012 ** 2, places=8)

    def test_kayit(self):
        with tempfile.TemporaryDirectory() as d:
            st = Store(Path(d) / "t.db")
            ref_id = st.referans_ekle("Proje", "#000000", self.fc, referans.sinir_katmani_oner(self.fc))
            self.assertEqual(st.referans(ref_id)["hesap_katmani"], "SNR_OSB")
            st.referans_guncelle(ref_id, hesap_katmani="*")
            self.assertIsNone(st.referans(ref_id)["hesap_katmani"])
            st._db.close()


class CizimTest(unittest.TestCase):
    def test_cizim(self):
        fc = referans.cizimden([[31.40, 41.20], [31.41, 41.20], [31.41, 41.21]], "Deneme")
        halka = fc["features"][0]["geometry"]["coordinates"][0]
        self.assertEqual(halka[0], halka[-1])
        self.assertEqual(fc["features"][0]["properties"], {"ad": "Deneme", "katman": "Çizim"})

    def test_gecersiz_cizim(self):
        with self.assertRaises(ValueError):
            referans.cizimden([[31.40, 41.20], [31.41, 41.20]])
        with self.assertRaisesRegex(ValueError, "kendini kesiyor"):
            referans.cizimden([[0, 0], [1, 1], [1, 0], [0, 1]])  # papyon


if __name__ == "__main__":
    unittest.main()
