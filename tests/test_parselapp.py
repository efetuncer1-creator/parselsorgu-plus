"""Testler; TKGM'ye bağlanmaz. Çalıştırma:  python -m unittest -v"""

import io
import json
import re
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from parselapp import analiz, coords, exporters, guncelleme, referans, tablo
from parselapp.importers import dosya_oku, ice_aktar, kml_oku
from parselapp.store import Store
from parselapp.tkgm import Bulunamadi, LimitAsildi, Tkgm
from parselapp.util import parcel_key, parse_alan

ORNEK = Path(__file__).parent / "ornekler"


def oku(ad):
    return (ORNEK / ad).read_bytes()


class GeciciStore(unittest.TestCase):
    def setUp(self):
        self.klasor = tempfile.TemporaryDirectory()
        self.store = Store(Path(self.klasor.name) / "t.db")

    def tearDown(self):
        self.store._db.close()
        self.klasor.cleanup()


class YardimciTest(unittest.TestCase):
    def test_alan(self):
        self.assertEqual(parse_alan("4.598,40"), 4598.40)
        self.assertEqual(parse_alan("3.158.004,18"), 3158004.18)
        self.assertEqual(parse_alan("500"), 500.0)
        self.assertIsNone(parse_alan(""))
        # Haritadan (koordinatla) sorgu İngilizce biçim döndürüyor
        self.assertEqual(parse_alan("29,636.19"), 29636.19)
        self.assertEqual(parse_alan("1,234,567.5"), 1234567.5)
        self.assertEqual(parse_alan("578,84"), 578.84)
        self.assertEqual(parse_alan("578.84"), 578.84)
        self.assertEqual(parse_alan("3.158.004"), 3158004.0)

    def test_turkce_siralama(self):
        from parselapp.util import tr_sira
        iller = ["Zonguldak", "Çankırı", "İzmir", "Ağrı", "Isparta", "Şanlıurfa", "Ordu", "Ödemiş", "Uşak", "Ünye", "Cide"]
        self.assertEqual(sorted(iller, key=tr_sira),
                         ["Ağrı", "Cide", "Çankırı", "Isparta", "İzmir", "Ordu", "Ödemiş", "Şanlıurfa", "Uşak", "Ünye", "Zonguldak"])
        self.assertEqual(sorted(["Apasaraycık Köyü", "Apa Köyü"], key=tr_sira), ["Apa Köyü", "Apasaraycık Köyü"])

    def test_anahtar_turkce(self):
        self.assertEqual(parcel_key("İSTANBUL", "Büyükçekmece", "  Türkoba ", 0, "53"),
                         parcel_key("istanbul", "büyükçekmece", "türkoba", "0", 53))
        # TKGM bazen ı'yı i'ye çevirerek veriyor; aynı parsel tek anahtar
        self.assertEqual(parcel_key("Ağrı", "Doğubayazıt", "Atabakan", 1, 2),
                         parcel_key("Ağri", "Doğubayazit", "ATABAKAN", 1, 2))

    def test_iyi_yazim(self):
        from parselapp.util import iyi_yazim
        self.assertEqual(iyi_yazim("Kirşehir", "Kırşehir"), "Kırşehir")
        self.assertEqual(iyi_yazim("Kırşehir", "Kirşehir"), "Kırşehir")
        self.assertEqual(iyi_yazim("Ilgaz", "İlgaz"), "Ilgaz")
        self.assertEqual(iyi_yazim("Akçaali", "Akçaali-Merkez"), "Akçaali")  # başka ad: dokunma
        self.assertEqual(iyi_yazim("Ankara", None), "Ankara")
        self.assertEqual(iyi_yazim(None, "Ankara"), "Ankara")


class IceAktarmaTest(GeciciStore):
    def test_alan_yoneticisi_kml(self):
        kayitlar = kml_oku(oku("alan_yoneticisi.kml"))
        self.assertEqual(len(kayitlar), 2)  # etiket noktası atlanır
        k = kayitlar[0]
        self.assertEqual((k["il"], k["ilce"], k["mahalle"], k["ada"], k["parsel"]),
                         ("Örnekil", "Örnekilçe", "Deneme", "12", "5"))
        self.assertEqual(k["alan"], "1.234,50")
        self.assertEqual(k["nitelik"], "Tarla")
        self.assertEqual(len(k["geometri"][0][0]), 5)
        self.assertEqual(kayitlar[1]["ada"], "0")  # boş ada -> 0

    def test_parselsorgu_kml_kayan_nokta_temizlenir(self):
        k = kml_oku(oku("parselsorgu.kml"))[0]
        self.assertEqual(k["geometri"][0][0][0], [31.4, 41.2])
        self.assertEqual(k["parsel"], "5")

    def test_fav_list_json(self):
        k = dosya_oku("fav_list.json", oku("fav_list.json"))[0]
        self.assertEqual(k["mahalle_id"], 100)
        self.assertEqual(k["geometri"][0][0][1], [31.4014, 41.201])
        self.assertEqual(k["zemin"], "Ana Taşınmaz")

    def test_tekrarlar_ve_tarih_kurali(self):
        r1 = ice_aktar(self.store, "alan_yoneticisi.kml", oku("alan_yoneticisi.kml"), "2025-01-01T00:00:00")
        self.assertEqual({k: r1[k] for k in ("yeni", "guncellendi", "atlandi")}, {"yeni": 2, "guncellendi": 0, "atlandi": 0})
        self.assertEqual(len(r1["keys"]), 2)
        # Daha eski tarihli dosya yenisini ezmez
        r2 = ice_aktar(self.store, "parselsorgu.kml", oku("parselsorgu.kml"), "2024-01-01T00:00:00")
        self.assertEqual(r2["atlandi"], 1)
        self.assertEqual(len(r2["keys"]), 1)  # atlanan parsel de oturuma eklenir
        # Daha yeni tarihli dosya günceller
        r3 = ice_aktar(self.store, "parselsorgu.kml", oku("parselsorgu.kml"), "2026-01-01T00:00:00")
        self.assertEqual(r3["guncellendi"], 1)
        p = self.store.bul(parcel_key("Örnekil", "Örnekilçe", "Deneme", "12", "5"))
        self.assertEqual(p["nitelik"], "Bahçe")
        self.assertEqual(p["pafta"], "X00a00a0a")  # yeni dosyada olmayan bilgi korunur
        self.assertEqual(len(self.store.listele()), 2)

    def test_eski_dosya_bos_alanlari_doldurur(self):
        kayit = dosya_oku("fav_list.json", oku("fav_list.json"))[0]
        kml_hali = {k: v for k, v in kayit.items() if k not in ("mahalle_id", "zemin")}
        self.store.kaydet(dict(kml_hali, nitelik="Yeni"), "dosya:yeni.kml", "2026-01-01T00:00:00")
        self.assertEqual(self.store.kaydet(kayit, "dosya:fav_list.json", "2025-01-01T00:00:00"), "atlandi")
        p = self.store.bul_id(100, "12", "7")
        self.assertEqual(p["nitelik"], "Yeni")  # ezilmedi
        self.assertEqual(p["zemin"], "Ana Taşınmaz")  # ama boşluk dolduruldu

    def test_turkce_yazim_farki_tek_kayit(self):
        kayit = {"il": "Ağri", "ilce": "Doğubayazit", "mahalle": "Atabakan", "ada": "1", "parsel": "2",
                 "alan": "100", "geometri": [[[[30, 40], [30.001, 40], [30.001, 40.001], [30, 40]]]]}
        self.store.kaydet(kayit, "tkgm")
        self.assertEqual(self.store.kaydet(kayit | {"il": "Ağrı", "ilce": "Doğubayazıt"}, "dosya:a.kml"), "atlandi")
        hepsi = self.store.listele()
        self.assertEqual(len(hepsi), 1)
        self.assertEqual((hepsi[0]["il"], hepsi[0]["ilce"]), ("Ağrı", "Doğubayazıt"))
        self.assertEqual(hepsi[0]["kaynak"], "tkgm")

    def test_ayni_yerin_yazimlari_birlesir(self):
        geo = [[[[30, 40], [30.001, 40], [30.001, 40.001], [30, 40]]]]
        self.store.kaydet({"il": "Ağri", "ilce": "Doğubayazit", "mahalle": "Atabakan", "ada": "1", "parsel": "1",
                           "geometri": geo}, "dosya:eski.kml")
        ice_aktar_sonrasi = {"il": "Ağrı", "ilce": "Doğubayazıt", "mahalle": "Atabakan", "ada": "1", "parsel": "2",
                             "geometri": geo}
        self.store.kaydet(ice_aktar_sonrasi, "dosya:yeni.kml")
        self.store.yazimlari_birlestir()
        self.assertEqual({(p["il"], p["ilce"]) for p in self.store.listele()}, {("Ağrı", "Doğubayazıt")})

    def test_eski_anahtarlar_yenilenir(self):
        import sqlite3
        db = Path(self.klasor.name) / "eski.db"
        Store(db)._db.close()
        c = sqlite3.connect(db)
        geo = json.dumps([[[[30, 40], [30.001, 40], [30.001, 40.001], [30, 40]]]])
        for key, il, kaynak, tarih in (("ağri|x|y|1|2", "Ağri", "tkgm", "2026-10-01"),
                                       ("ağrı|x|y|1|2", "Ağrı", "dosya:a.kml", "2026-10-05")):
            c.execute("INSERT INTO parseller (key, il, ilce, mahalle, ada, parsel, mahalle_id, geometri, kaynak,"
                      " kaynak_tarih, eklenme) VALUES (?,?,'X','Y','1','2',?,?,?,?,?)",
                      (key, il, 7 if kaynak == "tkgm" else None, geo, kaynak, tarih, tarih))
        c.execute("PRAGMA user_version = 0"); c.commit(); c.close()
        s = Store(db)
        hepsi = s.listele()
        s._db.close()
        self.assertEqual(len(hepsi), 1)
        p = hepsi[0]
        self.assertEqual((p["key"], p["il"], p["kaynak"], p["mahalle_id"]), ("ağri|x|y|1|2", "Ağrı", "tkgm", 7))

    def test_tkgm_verisi_dosyayla_ezilmez(self):
        kayit = dosya_oku("fav_list.json", oku("fav_list.json"))[0]
        self.store.kaydet(kayit, "tkgm", "2020-01-01T00:00:00")
        kayit2 = dict(kayit, nitelik="Değişmiş")
        self.assertEqual(self.store.kaydet(kayit2, "dosya:x.kml", "2030-01-01T00:00:00"), "atlandi")

    def test_bos_secim_hepsi_demek_degil(self):
        ice_aktar(self.store, "alan_yoneticisi.kml", oku("alan_yoneticisi.kml"))
        self.assertEqual(self.store.listele([]), [])
        self.assertEqual(len(self.store.listele()), 2)
        self.assertEqual(self.store.sil([]), 0)

    def test_noktadaki_parsel(self):
        ice_aktar(self.store, "alan_yoneticisi.kml", oku("alan_yoneticisi.kml"))
        self.assertEqual(self.store.noktadaki(31.4002, 41.20015)["parsel"], "5")
        self.assertIsNone(self.store.noktadaki(31.5, 41.5))


class ProjeTest(GeciciStore):
    """Kaydedilen proje başka bir bilgisayarda (boş kayıtta) açılır."""

    def setUp(self):
        super().setUp()
        ice_aktar(self.store, "alan_yoneticisi.kml", oku("alan_yoneticisi.kml"))
        from parselapp.server import ozellik
        self.parseller = [ozellik(p) for p in self.store.listele()]
        fc = referans.dosya_oku("osb_sinir.kml", oku("osb_sinir.kml"))
        self.alan = {"ad": "Deneme OSB", "renk": "#FF9F0A", "geojson": fc, "gizli": ["PL"], "hesap_katmani": None,
                     "gorsel": {"desen": "kareli", "kalinlik": 5}, "gorunur": False, "taranan": [[31.4, 41.2, 12.5]]}

    def _uygulama(self):
        from parselapp.server import Uygulama
        k = tempfile.TemporaryDirectory()
        self.addCleanup(k.cleanup)
        bos = Store(Path(k.name) / "bos.db")
        self.addCleanup(bos._db.close)
        return Uygulama(bos)

    def test_yukle(self):
        uyg = self._uygulama()
        govde = json.dumps({"parseller": self.parseller, "alanlar": [self.alan]}).encode()
        r = uyg.yonlendir("POST", ["proje", "yukle"], govde, {})
        self.assertEqual(sorted(r["keys"]), sorted(p["properties"]["key"] for p in self.parseller))
        geri = uyg.store.listele(r["keys"])
        self.assertEqual({(p["ada"], p["parsel"], p["kaynak"]) for p in geri},
                         {(p["properties"]["ada"], p["properties"]["parsel"], p["properties"]["kaynak"]) for p in self.parseller})
        a = r["alanlar"][0]
        self.assertEqual((a["ad"], a["renk"], a["gorunur"], a["gizli"], a["taranan"], a["gorsel"]["desen"], a["gorsel"]["kalinlik"]),
                         ("Deneme OSB", "#ff9f0a", False, ["PL"], [[31.4, 41.2, 12.5]], "kareli", 5))

    def test_bozuk_proje_yari_yuklenmez(self):
        from parselapp.server import Hata
        uyg = self._uygulama()
        bozuklar = [
            {"parseller": self.parseller, "alanlar": [dict(self.alan, renk="kırmızı")]},
            {"parseller": self.parseller + [{"properties": {"ada": "1"}, "geometry": None}], "alanlar": []},
            {"parseller": [], "alanlar": [dict(self.alan, geojson={"type": "Point"})]},
            {"parseller": "yok", "alanlar": []},
        ]
        for govde in bozuklar:
            with self.subTest(govde=str(govde)[:60]), self.assertRaises(Hata):
                uyg.yonlendir("POST", ["proje", "yukle"], json.dumps(govde).encode(), {})
        self.assertEqual(uyg.store.listele(), [])
        self.assertEqual(uyg.store.referanslar(), [])

    def test_tkgm_verisi_eski_projeyle_ezilmez(self):
        uyg = self._uygulama()
        p = json.loads(json.dumps(self.parseller[0]))
        p["properties"].update(kaynak="tkgm", kaynak_tarih="2026-01-01T00:00:00", nitelik="Eski")
        uyg.yonlendir("POST", ["proje", "yukle"], json.dumps({"parseller": [p], "alanlar": []}).encode(), {})
        p["properties"].update(kaynak_tarih="2025-01-01T00:00:00", nitelik="Daha eski")
        uyg.yonlendir("POST", ["proje", "yukle"], json.dumps({"parseller": [p], "alanlar": []}).encode(), {})
        self.assertEqual(uyg.store.listele()[0]["nitelik"], "Eski")


class AnalizTest(GeciciStore):
    def test_alan_ici_ucu(self):
        from parselapp.server import Uygulama
        ice_aktar(self.store, "alan_yoneticisi.kml", oku("alan_yoneticisi.kml"))
        uyg = Uygulama(self.store)
        ref_id = str(self.store.referans_ekle("OSB", "#bf5af2", referans.dosya_oku("osb_sinir.kml", oku("osb_sinir.kml"))))
        anahtar = {p["parsel"]: p["key"] for p in self.store.listele()}
        r = uyg.yonlendir("POST", ["referanslar", ref_id, "alan-ici"], json.dumps({"keys": list(anahtar.values())}).encode(), {})
        self.assertEqual(set(r), {anahtar["5"], anahtar["1"]})  # 12/7 alana değmiyor
        m2, oran = r[anahtar["5"]]
        self.assertGreater(m2, 0)
        self.assertGreaterEqual(oran, 0.98)
        icinde = uyg.yonlendir("POST", ["referanslar", ref_id, "alan-ici"],
                               json.dumps({"keys": list(anahtar.values()), "kapsam": "icinde"}).encode(), {})
        self.assertEqual(set(icinde), {anahtar["5"]})  # 13/1 alana kısmen giriyor

    def test_tablo_excel(self):
        from parselapp.server import Uygulama
        uyg = Uygulama(self.store)
        sayfalar = [
            {"ad": "Parsel türleri", "sutunlar": [{"ad": "Nitelik", "tur": "metin"}, {"ad": "Adet", "tur": "tamsayi"},
                                                  {"ad": "Alan (m²)", "tur": "sayi"}, {"ad": "Oran", "tur": "yuzde"}],
             "satirlar": [["Tarla", 3, 12683.71, 0.75], ["Bahçe", 1, 4200.5, 0.25]], "toplam": ["Toplam", 4, 16884.21, 1]},
            {"ad": "Parseller: 1/2", "sutunlar": [{"ad": "Ada", "tur": "kod"}, {"ad": "Parsel", "tur": "kod"}],
             "satirlar": [["512", "15"], ["513", "2a"]]},
        ]
        dosya = uyg.yonlendir("POST", ["tablo-excel"], json.dumps({"sayfalar": sayfalar, "dosyaAdi": "Çumra/analiz"}).encode(), {})
        self.assertEqual(dosya.ad, "Çumra-analiz.xlsx")
        kitap = load_workbook(io.BytesIO(dosya.veri))
        self.assertEqual(kitap.sheetnames, ["Parsel türleri", "Parseller- 1-2"])
        s = kitap["Parsel türleri"]
        self.assertEqual([c.value for c in s[2]], ["Tarla", 3, 12683.71, 0.75])
        self.assertEqual(s["C2"].number_format, "#,##0.00")
        self.assertEqual(s["D2"].number_format, "0.0%")
        self.assertTrue(s["A4"].font.bold)
        self.assertEqual([c.value for c in kitap["Parseller- 1-2"][3]], [513, "2a"])  # sayı olan kod sayıya döner

    def test_tablo_excel_bozuk(self):
        from parselapp.server import Hata, Uygulama
        uyg = Uygulama(self.store)
        for sayfalar in (None, [], [{"ad": "x", "sutunlar": [{"ad": "a", "tur": "renk"}], "satirlar": []}],
                         [{"ad": "x", "sutunlar": [{"ad": "a", "tur": "metin"}], "satirlar": [["1", "2"]]}]):
            with self.subTest(sayfalar=sayfalar), self.assertRaises(Hata):
                uyg.yonlendir("POST", ["tablo-excel"], json.dumps({"sayfalar": sayfalar}).encode(), {})


class IzgaraTest(unittest.TestCase):
    def test_izgara(self):
        # Çumra çevresinde ~1,2 km x 1,5 km
        iz = coords.izgara(32.798, 37.585, 32.812, 37.599)
        self.assertEqual((iz["sistem"], iz["dom"]), ("ITRF96 3°", 33))
        self.assertIn(iz["aralik"], coords.IZGARA_ARALIKLARI)
        self.assertTrue(1 <= len(iz["yler"]) <= 7 and 1 <= len(iz["xler"]) <= 7)
        for cizgi in iz["yler"] + iz["xler"]:
            self.assertEqual(cizgi["deger"] % iz["aralik"], 0)
            for lon, lat in cizgi["noktalar"]:
                self.assertTrue(32.79 < lon < 32.82 and 37.58 < lat < 37.61)
        # Y çizgisi boyunca dilim düzleminde Y sabit kalır
        y = iz["yler"][0]
        for lon, lat in y["noktalar"][::4]:
            self.assertAlmostEqual(coords.donustur(lon, lat, "itrf96", 33)[0], y["deger"], delta=0.01)

    def test_izgara_ucu(self):
        from parselapp.server import Hata, Uygulama
        with tempfile.TemporaryDirectory() as k:
            uyg = Uygulama(Store(Path(k) / "t.db"))
            iz = uyg.yonlendir("POST", ["izgara"], json.dumps({"bbox": [32.798, 37.585, 32.812, 37.599], "sistem": "ed50"}).encode(), {})
            self.assertEqual(iz["sistem"], "ED50 3°")
            for bozuk in ({"bbox": [1, 2, 3, 4]}, {"bbox": [32.8, 37.6, 32.7, 37.5]}, {"bbox": [32.79, 37.58, 32.81, 37.6], "sistem": "utm"}, {}):
                with self.subTest(govde=bozuk), self.assertRaises(Hata):
                    uyg.yonlendir("POST", ["izgara"], json.dumps(bozuk).encode(), {})
            uyg.store._db.close()


class KoordinatTest(unittest.TestCase):
    def test_dilim(self):
        self.assertEqual(coords.dilim_sec(31.42), 30)
        self.assertEqual(coords.dilim_sec(31.6), 33)
        self.assertEqual(coords.dilim_sec(26.0), 27)
        self.assertEqual(coords.dilim_sec(44.9), 45)

    def test_itrf96_ed50(self):
        y, x = coords.donustur(30.0, 41.0, "itrf96", 30)
        self.assertAlmostEqual(y, 500000.0, delta=0.01)  # orta meridyende Y = 500.000
        self.assertAlmostEqual(x, 4540000, delta=10000)
        y2, x2 = coords.donustur(30.0, 41.0, "ed50", 30)
        # ED50 ile ITRF96 arası Türkiye'de birkaç yüz metre değil, ~70-200 m civarı fark
        self.assertTrue(20 < abs(x2 - x) < 300)

    def test_geri_donusum(self):
        for sistem in coords.SISTEMLER:
            y, x = coords.donustur(27.85, 41.08, sistem, 27)
            boylam, enlem = coords.cografyaya(y, x, sistem, 27)
            self.assertAlmostEqual(boylam, 27.85, places=7)
            self.assertAlmostEqual(enlem, 41.08, places=7)

    def test_koordinat_ucu(self):
        from parselapp.server import Hata, Uygulama
        with tempfile.TemporaryDirectory() as k:
            uyg = Uygulama(Store(Path(k) / "t.db"))
            y, x = coords.donustur(27.85, 41.08, "ed50", 27)
            r = uyg.yonlendir("POST", ["koordinat"], json.dumps({"y": y, "x": x, "dilim": 27}).encode(), {})
            self.assertEqual([s["sistem"] for s in r["sonuclar"]], ["itrf96", "ed50"])
            ed50 = r["sonuclar"][1]
            self.assertAlmostEqual(ed50["boylam"], 27.85, places=6)
            self.assertAlmostEqual(ed50["enlem"], 41.08, places=6)
            for bozuk in ({"y": y, "x": x, "dilim": 28}, {"y": 4550000, "x": 567000, "dilim": 27}, {"y": y, "x": x}, {}):
                with self.subTest(govde=bozuk), self.assertRaises(Hata):
                    uyg.yonlendir("POST", ["koordinat"], json.dumps(bozuk).encode(), {})
            uyg.store._db.close()


class DisaAktarmaTest(GeciciStore):
    def setUp(self):
        super().setUp()
        ice_aktar(self.store, "alan_yoneticisi.kml", oku("alan_yoneticisi.kml"))
        ice_aktar(self.store, "fav_list.json", oku("fav_list.json"))
        self.parseller = self.store.listele()

    def test_kml_geri_okunur(self):
        geri = kml_oku(exporters.kml(self.parseller))
        self.assertEqual(len(geri), 3)
        anahtarlar = {parcel_key(k["il"], k["ilce"], k["mahalle"], k["ada"], k["parsel"]) for k in geri}
        self.assertEqual(anahtarlar, {p["key"] for p in self.parseller})
        fav = next(k for k in geri if k["parsel"] == "7")
        self.assertEqual(fav["mahalle_id"], 100)

    def test_kml_ayri_zip(self):
        z = zipfile.ZipFile(io.BytesIO(exporters.kml_ayri_zip(self.parseller)))
        self.assertEqual(len(z.namelist()), 3)
        self.assertIn("Örnekil-Örnekilçe-Deneme-12-5.kml", z.namelist())

    def test_excel(self):
        veri, dom = exporters.excel(self.parseller, "itrf96")
        self.assertEqual(dom, 30)
        wb = load_workbook(io.BytesIO(veri))
        ws = wb["Parseller"]
        self.assertEqual(ws["G1"].value, "Tapu Alanı (m²)")
        alanlar = [ws.cell(row=r, column=7).value for r in range(2, 5)]
        self.assertEqual(sorted(alanlar), [500.0, 1234.5, 2000.0])
        self.assertEqual(ws["G5"].value, "=SUM(G2:G4)")
        wk = wb["Koordinatlar"]
        self.assertIn("EPSG:5254", wk["A1"].value)
        # 12/5 dış halkası 5 nokta, kapanış noktası tekrarlanmaz -> 4 satır
        satirlar = [r for r in wk.iter_rows(min_row=5, values_only=True) if r[1] == "12/5"]
        self.assertEqual(len(satirlar), 4)

    def test_excel_ed50_sabit_dilim(self):
        _, dom = exporters.excel(self.parseller, "ed50", "33")
        self.assertEqual(dom, 33)
        with self.assertRaises(ValueError):
            exporters.excel(self.parseller, "ed50", "31")


class ReferansTest(GeciciStore):
    def setUp(self):
        super().setUp()
        ice_aktar(self.store, "alan_yoneticisi.kml", oku("alan_yoneticisi.kml"))
        ice_aktar(self.store, "fav_list.json", oku("fav_list.json"))
        self.fc = referans.dosya_oku("osb_sinir.kml", oku("osb_sinir.kml"))

    def test_okuma(self):
        turler = sorted(f["geometry"]["type"] for f in self.fc["features"])
        self.assertEqual(turler, ["LineString", "Point", "Polygon"])
        self.assertEqual(self.fc["features"][0]["properties"]["ad"], "1. Etap")

    def test_kmz(self):
        tampon = io.BytesIO()
        with zipfile.ZipFile(tampon, "w") as z:
            z.writestr("doc.kml", oku("osb_sinir.kml"))
        fc = referans.dosya_oku("osb.kmz", tampon.getvalue())
        self.assertEqual(len(fc["features"]), 3)

    def test_kapali_cizgi_alan_sayilir(self):
        alan = referans.alan(self.fc)
        self.assertEqual(alan.geom_type, "MultiPolygon")  # poligon + çizgiden oluşan alan
        self.assertEqual(len(alan.geoms), 2)

    def test_icindeki_ve_kesisen_parseller(self):
        parseller = self.store.listele()
        anahtar = {p["parsel"]: p["key"] for p in parseller}
        o = analiz.ozet({"geojson": self.fc, "gizli": [], "taranan": []}, parseller)
        icinde, kesisen = o["icinde"]["keys"], o["kesisen"]["keys"]
        self.assertEqual(icinde, [anahtar["5"]])                     # 12/5 tamamen 1. etapta
        self.assertEqual(sorted(kesisen), sorted([anahtar["5"], anahtar["1"]]))  # 13/1 2. etaba kısmen giriyor
        self.assertNotIn(anahtar["7"], kesisen)                      # 12/7 dışarıda

    def test_katmanlar(self):
        self.assertEqual(referans.katmanlar(self.fc), [("Deneme OSB", 3)])
        kml = (b'<kml xmlns="http://www.opengis.net/kml/2.2"><Document><name>Proje</name>'
               b'<Folder><name>KADASTRO</name><Placemark><Polygon><outerBoundaryIs><LinearRing><coordinates>'
               b'31.3999,41.1999 31.40045,41.1999 31.40045,41.20035 31.3999,41.1999</coordinates></LinearRing>'
               b'</outerBoundaryIs></Polygon></Placemark></Folder>'
               b'<Folder><name>PL</name><Folder><name>ALAN</name><Placemark><Polygon><outerBoundaryIs><LinearRing>'
               b'<coordinates>31.4005,41.1999 31.4008,41.1999 31.4008,41.2002 31.4005,41.1999</coordinates>'
               b'</LinearRing></outerBoundaryIs></Polygon></Placemark></Folder></Folder></Document></kml>')
        fc = referans.dosya_oku("proje.kml", kml)
        self.assertEqual(referans.katmanlar(fc), [("KADASTRO", 1), ("PL / ALAN", 1)])
        self.assertEqual(len(referans.alan(fc).geoms), 2)
        self.assertEqual(referans.alan(fc, ["PL / ALAN"]).geom_type, "Polygon")  # gizli katman hesaba girmez
        self.assertIsNone(referans.alan(fc, ["KADASTRO", "PL / ALAN"]))

    def test_alani_olmayan_referans(self):
        fc = {"type": "FeatureCollection", "features": [
            {"type": "Feature", "geometry": {"type": "Point", "coordinates": [31.4, 41.2]}, "properties": {}}]}
        self.assertIsNone(referans.alan(fc))
        with self.assertRaises(ValueError):
            analiz.ozet({"geojson": fc, "gizli": [], "taranan": []}, self.store.listele())

    def test_kayit(self):
        ref_id = self.store.referans_ekle("Deneme OSB", "#2b7bff", self.fc)
        self.store.referans_guncelle(ref_id, gorunur=0, renk="#ff0000", gizli=["PL", "PL"], baska="yok sayılır")
        r = self.store.referans(ref_id)
        self.assertEqual((r["ad"], r["renk"], r["gorunur"], r["gizli"]), ("Deneme OSB", "#ff0000", False, ["PL"]))
        self.assertEqual(len(r["geojson"]["features"]), 3)
        self.assertEqual(self.store.referans_sil(ref_id), 1)
        self.assertEqual(self.store.referanslar(), [])

    def test_gorsel_varsayilan_ve_kayit(self):
        ref_id = self.store.referans_ekle("Deneme OSB", "#2b7bff", self.fc)
        self.assertEqual(self.store.referans(ref_id)["gorsel"], referans.GORSEL)
        g = referans.gorsel_duzelt({"kalinlik": 3.3, "desen": "kareli", "aralik": 100, "opaklik": 0, "etiket": False})
        self.store.referans_guncelle(ref_id, gorsel=g)
        self.assertEqual(self.store.referans(ref_id)["gorsel"],
                         {"kalinlik": 3.5, "cizgi": "kesik", "desen": "kareli", "aralik": 24, "opaklik": 5, "etiket": False,
                          "katman_opaklik": 100})

    def test_gorsel_dogrulama(self):
        for bozuk in ({"desen": "zigzag"}, {"cizgi": None}, {"kalinlik": "3"}, {"kalinlik": True},
                      {"aralik": float("nan")}, {"etiket": "evet"}, ["kalinlik"]):
            with self.subTest(bozuk=bozuk), self.assertRaises(ValueError):
                referans.gorsel_duzelt(bozuk)
        self.assertEqual(referans.renk_duzelt("#AABBCC"), "#aabbcc")
        for bozuk in ("red", "#abc", "#aabbcc; x", None):
            with self.subTest(renk=bozuk), self.assertRaises(ValueError):
                referans.renk_duzelt(bozuk)
        self.assertEqual(referans.gorsel_oku("bozuk json"), referans.GORSEL)

    def test_eski_veritabanina_gorsel_eklenir(self):
        import sqlite3
        db = Path(self.klasor.name) / "eski.db"
        c = sqlite3.connect(db)
        c.execute("CREATE TABLE referanslar (id INTEGER PRIMARY KEY AUTOINCREMENT, ad TEXT NOT NULL, renk TEXT NOT NULL,"
                  " gorunur INTEGER NOT NULL DEFAULT 1, geojson TEXT NOT NULL, eklenme TEXT NOT NULL)")
        c.execute("INSERT INTO referanslar (ad, renk, geojson, eklenme) VALUES ('Eski', '#bf5af2', ?, '2026-10-01')",
                  (json.dumps(self.fc),))
        c.execute("PRAGMA user_version = 1"); c.commit(); c.close()
        s = Store(db)
        r = s.referanslar()[0]
        s._db.close()
        self.assertEqual((r["ad"], r["gorsel"]), ("Eski", referans.GORSEL))

    def test_guncelle_ucu(self):
        from parselapp.server import Hata, Uygulama
        uyg = Uygulama(self.store)
        ref_id = str(self.store.referans_ekle("Deneme OSB", "#2b7bff", self.fc))
        govde = json.dumps({"renk": "#FF9F0A", "gorsel": {"desen": "nokta", "kalinlik": 5}}).encode()
        r = uyg.yonlendir("POST", ["referanslar", ref_id, "guncelle"], govde, {})
        self.assertEqual((r["renk"], r["gorsel"]["desen"], r["gorsel"]["kalinlik"], r["gorsel"]["cizgi"]),
                         ("#ff9f0a", "nokta", 5, "kesik"))
        for bozuk in ({"renk": "kırmızı"}, {"gorsel": {"desen": "yok-boyle"}}):
            with self.subTest(govde=bozuk), self.assertRaises(Hata) as h:
                uyg.yonlendir("POST", ["referanslar", ref_id, "guncelle"], json.dumps(bozuk).encode(), {})
            self.assertEqual(h.exception.durum, 400)
        self.assertEqual(self.store.referans(int(ref_id))["gorsel"]["desen"], "nokta")  # hatalı istek kaydı bozmaz

    def test_bos_dosya_hata_verir(self):
        with self.assertRaises(ValueError):
            referans.dosya_oku("bos.kml", b'<kml xmlns="http://www.opengis.net/kml/2.2"><Document/></kml>')


class _Cevap:
    def __init__(self, durum, govde):
        self.status_code = durum
        self.text = govde if isinstance(govde, str) else json.dumps(govde)

    def json(self):
        return json.loads(self.text)


class _SahteOturum:
    def __init__(self, cevaplar):
        self.cevaplar = list(cevaplar)
        self.istekler = []
        self.headers = {}

    def get(self, adres, timeout=None):
        self.istekler.append(adres)
        return self.cevaplar.pop(0)


class TkgmTest(GeciciStore):
    def setUp(self):
        super().setUp()
        import parselapp.tkgm as t
        self._bekleme, t.BEKLEME_SN = t.BEKLEME_SN, 0

    def tearDown(self):
        import parselapp.tkgm as t
        t.BEKLEME_SN = self._bekleme
        super().tearDown()

    def _parsel_cevabi(self):
        return json.loads(oku("fav_list.json"))["features"][0] | {
            "geometry": {"type": "Polygon", "coordinates": [[[31.401, 41.201], [31.4014, 41.201], [31.4014, 41.2013], [31.401, 41.201]]]}}

    def test_parsel_bir_kez_sorgulanir(self):
        oturum = _SahteOturum([_Cevap(200, self._parsel_cevabi())])
        t = Tkgm(self.store, oturum)
        p, yeni = t.parsel(100, "12", "7")
        self.assertTrue(yeni)
        self.assertEqual(p["kaynak"], "tkgm")
        p2, yeni2 = t.parsel(100, "12", "7")
        self.assertFalse(yeni2)
        self.assertEqual(len(oturum.istekler), 1)
        self.assertTrue(oturum.istekler[0].endswith("/parsel/100/12/7"))
        self.assertEqual(self.store.bugunku_sorgu()["toplam"], 1)

    def test_nokta_sorgusu_kayittaysa_istek_atmaz(self):
        oturum = _SahteOturum([_Cevap(200, self._parsel_cevabi())])
        t = Tkgm(self.store, oturum)
        t.parsel(100, 12, 7)
        p, yeni = t.parsel_noktadan(41.2011, 31.4012)
        self.assertFalse(yeni)
        self.assertEqual(p["parsel"], "7")
        self.assertEqual(len(oturum.istekler), 1)

    def test_limit(self):
        t = Tkgm(self.store, _SahteOturum([_Cevap(403, '"Günlük sorgu limitini aştınız."')]))
        with self.assertRaises(LimitAsildi):
            t.parsel(100, 1, 1)
        self.assertEqual(self.store.bugunku_sorgu(), {"toplam": 1, "limit_hatasi": 1})

    def test_bulunamadi(self):
        t = Tkgm(self.store, _SahteOturum([_Cevap(200, "null")]))
        with self.assertRaises(Bulunamadi):
            t.parsel(100, 1, 999)

    def test_adlar_idari_listeden_duzeltilir(self):
        cevap = self._parsel_cevabi()
        cevap["properties"] = cevap["properties"] | {"ilAd": "Kirşehir", "ilId": 62}
        self.store.idari_kaydet("il", 0, [{"id": 62, "ad": "Kırşehir"}])
        t = Tkgm(self.store, _SahteOturum([_Cevap(200, cevap)]))
        p, _ = t.parsel(100, 12, 7)
        self.assertEqual(p["il"], "Kırşehir")

    def test_liste_gelince_kayitli_adlar_duzelir(self):
        cevap = self._parsel_cevabi()
        cevap["properties"] = cevap["properties"] | {"mahalleAd": "Sinanli"}
        liste = {"type": "FeatureCollection", "features": [
            {"type": "Feature", "geometry": None, "properties": {"text": "Sinanlı", "id": 100}}]}
        t = Tkgm(self.store, _SahteOturum([_Cevap(200, cevap), _Cevap(200, liste)]))
        p, _ = t.parsel(100, 12, 7)
        self.assertEqual(p["mahalle"], "Sinanli")
        t.mahalleler(10)
        self.assertEqual(self.store.bul_id(100, "12", "7")["mahalle"], "Sinanlı")

    def test_idari_liste_kaydedilir(self):
        cevap = {"type": "FeatureCollection", "features": [
            {"type": "Feature", "geometry": None, "properties": {"text": "Zonguldak", "id": 89}},
            {"type": "Feature", "geometry": None, "properties": {"text": "Adana", "id": 23}}]}
        oturum = _SahteOturum([_Cevap(200, cevap)])
        t = Tkgm(self.store, oturum)
        self.assertEqual(t.iller(), [{"id": 23, "ad": "Adana"}, {"id": 89, "ad": "Zonguldak"}])
        self.assertEqual(t.iller()[0]["ad"], "Adana")
        self.assertEqual(len(oturum.istekler), 1)


if __name__ == "__main__":
    unittest.main()


class TabloTest(unittest.TestCase):
    def test_excel(self):
        from datetime import date
        from openpyxl import Workbook
        kitap = Workbook()
        s = kitap.active
        s.title = "Parseller"
        s.append(["Ada", "Parsel", "Alan (m²)", "Tarih", None])
        s.append([512, 15, 12683.71, date(2026, 10, 9), None])
        s.append([513, 2, 1250.0, None, None])
        s.append([None, None, None, None, None])
        s["C2"].number_format = s["C3"].number_format = "#,##0.00"
        kitap.create_sheet("Boş")
        tampon = io.BytesIO()
        kitap.save(tampon)
        sonuc = tablo.oku("liste.xlsx", tampon.getvalue())
        self.assertEqual([x["ad"] for x in sonuc["sayfalar"]], ["Parseller"])  # boş sayfa atlanır
        self.assertEqual(sonuc["sayfalar"][0]["satirlar"], [
            ["Ada", "Parsel", "Alan (m²)", "Tarih"],
            ["512", "15", "12.683,71", "09.10.2026"],
            ["513", "2", "1.250,00", ""],
        ])
        self.assertFalse(sonuc["sayfalar"][0]["kesildi"])

    def test_hucre_bicimi(self):
        # Sayılar Excel'de hücre biçimiyle nasıl görünüyorsa öyle (koordinatlar kırpılmasın)
        from openpyxl import Workbook
        kitap = Workbook()
        s = kitap.active
        ornekler = [
            (7000, "#,##0.00", "7.000,00"), (15227.36, "#,##0.00", "15.227,36"),
            (4422771.3, "0.00", "4422771,30"), (525312.84, "0.00", "525312,84"),
            (39.93876, "General", "39,93876"), (2009, "General", "2009"), (4423399.0, "General", "4423399"),
            (0.256, "0.0%", "25,6%"), (1.5, "0.0#", "1,5"), (1.234, "0.0#", "1,23"), (12.7, "0", "13"),
            (-3.25, "#,##0.00;[Red]-#,##0.00", "-3,25"), (1234.5, '#,##0.0 "m²"', "1.234,5"),
        ]
        for deger, bicim, _ in ornekler:
            s.append([deger])
            s.cell(s.max_row, 1).number_format = bicim
        tampon = io.BytesIO()
        kitap.save(tampon)
        satirlar = tablo.oku("bicim.xlsx", tampon.getvalue())["sayfalar"][0]["satirlar"]
        self.assertEqual([r[0] for r in satirlar], [beklenen for *_, beklenen in ornekler])

    def test_excel_gorunusu(self):
        # Kalın/renkli başlık, sayılar sağa, açık hizalama, sütun genişliği Excel'deki gibi gelir
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        kitap = Workbook()
        s = kitap.active
        s.append(["Ada", "Alan", "Not"])
        s.append([512, 1250.5, "orta"])
        for h in s[1]:
            h.font = Font(b=True, color="FFFFFFFF")
            h.fill = PatternFill("solid", fgColor="FF2F5D50")
        s["C2"].alignment = Alignment(horizontal="center")
        s.column_dimensions["A"].width = 6
        tampon = io.BytesIO()
        kitap.save(tampon)
        sayfa = tablo.oku("gorunus.xlsx", tampon.getvalue())["sayfalar"][0]
        stil = lambda i, j: sayfa["stiller"][sayfa["stil"][i][j]]
        self.assertEqual(stil(0, 0), {"k": 1, "r": "#ffffff", "d": "#2f5d50"})
        self.assertEqual(stil(1, 0), {"h": "r"})  # sayı sağa
        self.assertEqual(stil(1, 2), {"h": "c"})  # açık hizalama
        self.assertEqual(sayfa["genislikler"][0], 6)
        self.assertEqual(sayfa["genislikler"][1], 8.43)  # tanımsız: Excel'in varsayılanı

    def test_csv_noktali_virgul_ve_sinir(self):
        satirlar = ["Ada;Parsel"] + [f"{i};{i + 1}" for i in range(tablo.EN_FAZLA_SATIR + 5)]
        sonuc = tablo.oku("liste.csv", "\n".join(satirlar).encode("utf-8-sig"))
        sayfa = sonuc["sayfalar"][0]
        self.assertEqual(sayfa["satirlar"][:2], [["Ada", "Parsel"], ["0", "1"]])
        self.assertEqual(len(sayfa["satirlar"]), tablo.EN_FAZLA_SATIR)
        self.assertTrue(sayfa["kesildi"])

    def test_kullanilan_alani_yanlis_yazilmis_excel(self):
        # Bazı programlar <dimension ref="A1"/> yazar; satırların hepsi yine okunmalı.
        from openpyxl import Workbook
        kitap = Workbook()
        for satir in (["Ada", "Parsel"], [512, 15], [513, 2]):
            kitap.active.append(satir)
        tampon = io.BytesIO()
        kitap.save(tampon)
        bozuk = io.BytesIO()
        with zipfile.ZipFile(tampon) as z, zipfile.ZipFile(bozuk, "w") as y:
            for ad in z.namelist():
                veri = z.read(ad)
                if ad == "xl/worksheets/sheet1.xml":
                    # Yazılış openpyxl'in XML yazıcısına göre değişir (lxml: "/>", standart: " />")
                    veri, n = re.subn(rb'<dimension ref="A1:B3"\s*/>', b'<dimension ref="A1"/>', veri)
                    self.assertEqual(n, 1)
                y.writestr(ad, veri)
        satirlar = tablo.oku("baska-program.xlsx", bozuk.getvalue())["sayfalar"][0]["satirlar"]
        self.assertEqual(satirlar, [["Ada", "Parsel"], ["512", "15"], ["513", "2"]])

    def test_windows_1254_csv(self):
        sonuc = tablo.oku("excel.csv", "İl;İlçe\nKonya;Çumra".encode("cp1254"))
        self.assertEqual(sonuc["sayfalar"][0]["satirlar"], [["İl", "İlçe"], ["Konya", "Çumra"]])

    def test_desteklenmeyen(self):
        for ad, icerik in (("eski.xls", b"x"), ("resim.png", b"x"), ("bozuk.xlsx", b"PK bozuk")):
            with self.subTest(ad=ad), self.assertRaises(ValueError):
                tablo.oku(ad, icerik)


class GuncellemeTest(unittest.TestCase):
    """GitHub'a bağlanmaz: son sürüm yanıtı sahte."""

    @staticmethod
    def yanit(etiket, dosyalar=("ParselSorguPlus-Kurulum-0.2.0.exe",)):
        return lambda adres: {
            "tag_name": etiket, "name": "Görsel sekmesi", "html_url": "https://github.com/x/y/releases/tag/" + etiket,
            "assets": [{"name": d, "browser_download_url": "https://github.com/x/y/releases/download/" + d} for d in dosyalar],
        }

    def test_surum_parcalari(self):
        self.assertEqual(guncelleme.parcalar("v0.2.0"), (0, 2, 0))
        self.assertEqual(guncelleme.parcalar("0.10"), (0, 10, 0))
        self.assertIsNone(guncelleme.parcalar("gelistirme"))
        self.assertIsNone(guncelleme.parcalar(None))

    def test_yeni_surum(self):
        r = guncelleme.denetle("0.1.0", self.yanit("v0.2.0"))
        self.assertTrue(r["yeni"])
        self.assertEqual(r["son"], "0.2.0")
        self.assertTrue(r["indir"].endswith(".exe"))
        # 0.10 > 0.9 (metin olarak değil, sayı olarak karşılaştırılır)
        self.assertTrue(guncelleme.denetle("0.9.0", self.yanit("v0.10.0"))["yeni"])

    def test_guncel_ya_da_daha_yeni(self):
        self.assertFalse(guncelleme.denetle("0.2.0", self.yanit("v0.2.0"))["yeni"])
        self.assertFalse(guncelleme.denetle("0.3.0", self.yanit("v0.2.0"))["yeni"])

    def test_kurulum_dosyasi_yoksa_surum_sayfasi(self):
        r = guncelleme.denetle("0.1.0", self.yanit("v0.2.0", dosyalar=()))
        self.assertEqual(r["indir"], r["sayfa"])

    def test_denetlenemezse_uyari_yok(self):
        def kopuk(adres):
            raise OSError("internet yok")
        self.assertFalse(guncelleme.denetle("0.1.0", kopuk)["yeni"])
        self.assertFalse(guncelleme.denetle("0.1.0", lambda a: {"tag_name": "son"})["yeni"])

    def test_kaynaktan_calisirken_denetlenmez(self):
        cagri = []
        r = guncelleme.denetle(None, lambda a: cagri.append(a))
        self.assertFalse(r["yeni"])
        self.assertEqual(cagri, [])

    def test_denetci_bir_kez_sorar(self):
        cagri = []
        def getir(adres):
            cagri.append(adres)
            return self.yanit("v0.2.0")(adres)
        d = guncelleme.Denetci("0.1.0", getir).baslat().baslat()
        self.assertTrue(d.sonuc()["yeni"])
        self.assertTrue(d.sonuc()["yeni"])
        self.assertEqual(len(cagri), 1)

